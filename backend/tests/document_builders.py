"""Build small document files in memory, so the tests need no binary fixtures.

Shared by the parser and ingest tests. ``tests/`` is not a package: pytest puts
this directory on ``sys.path``, so tests import it as ``document_builders``.
"""
import io
import zipfile
from email.message import EmailMessage


def make_docx() -> bytes:
    from docx import Document

    document = Document()
    document.core_properties.title = "Rapport annuel"
    document.core_properties.author = "Marie Curie"
    document.core_properties.language = "fr-FR"
    document.sections[0].header.paragraphs[0].text = "EN-TÊTE RÉPÉTÉ"
    document.sections[0].footer.paragraphs[0].text = "PIED DE PAGE"
    document.add_heading("Chapitre 1", level=1)
    document.add_paragraph("Introduction du rapport.")
    document.add_paragraph("premier point", style="List Bullet")
    document.add_heading("Budget", level=2)
    table = document.add_table(rows=3, cols=2)
    for r, row in enumerate([["Poste", "Montant"], ["Loyer", "1200"], ["Eau", "30"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _odf_table(rows, name="T"):
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    table = Table(name=name)
    for row in rows:
        tr = TableRow()
        table.addElement(tr)
        for value in row:
            cell = TableCell()
            cell.addElement(P(text=value))
            tr.addElement(cell)
    return table


def make_odt() -> bytes:
    from odf import dc
    from odf.opendocument import OpenDocumentText
    from odf.text import H, List, ListItem, P

    document = OpenDocumentText()
    document.meta.addElement(dc.Title(text="Arrêté préfectoral"))
    document.meta.addElement(dc.Creator(text="Préfecture"))
    document.meta.addElement(dc.Language(text="fr-FR"))
    document.text.addElement(H(outlinelevel=1, text="Titre I"))
    document.text.addElement(H(outlinelevel=2, text="Dispositions générales"))
    document.text.addElement(P(text="Le présent arrêté s'applique."))
    items = List()
    item = ListItem()
    item.addElement(P(text="premier alinéa"))
    items.addElement(item)
    document.text.addElement(items)
    document.text.addElement(_odf_table([["Commune", "Population"], ["Lyon", "522000"]]))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_ods() -> bytes:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import TableCell, TableRow

    document = OpenDocumentSpreadsheet()
    table = _odf_table([["Poste", "Montant"], ["Loyer", "1200"]], name="Budget")
    # LibreOffice pads sheets with huge runs of empty repeated cells and rows.
    for row in table.childNodes:
        row.addElement(TableCell(numbercolumnsrepeated=16000))
    table.addElement(TableRow(numberrowsrepeated=1_000_000))
    document.spreadsheet.addElement(table)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.core_properties.title = "Bilan"
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Introduction"
    slide.placeholders[1].text = "Objectifs atteints"
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Chiffres"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2)).table
    for (r, c), value in {(0, 0): "Clé", (0, 1): "Valeur", (1, 0): "x", (1, 1): "1"}.items():
        table.cell(r, c).text = value
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def make_pdf(pages: int = 1, password: str | None = None) -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(200, 200)
    if password:
        writer.encrypt(password)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_eml(attachments=()) -> bytes:
    message = EmailMessage()
    message["From"] = "Alice Martin <alice@example.fr>"
    message["To"] = "bob@example.fr"
    message["Subject"] = "Compte rendu de réunion"
    message["Date"] = "Mon, 02 Mar 2026 10:00:00 +0100"
    message.set_content("Bonjour,\n\nVoici le compte rendu.")
    message.add_alternative("<p>Version <b>HTML</b></p>", subtype="html")
    for name, payload, maintype, subtype in attachments:
        message.add_attachment(payload, maintype=maintype, subtype=subtype, filename=name)
    # Set last: building the body moves Content-* headers onto the first part.
    message["Content-Language"] = "fr"
    return message.as_bytes()


def zip_bytes(members: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    return buffer.getvalue()
