"""PowerPoint .pptx parser (python-pptx): one section per slide."""
from __future__ import annotations

from io import BytesIO
from typing import Any

from app.rag.parsers.archive import open_zip
from app.rag.parsers.base import (
    DocumentMetadata,
    ParsedDocument,
    ParseError,
    clean_str,
    iso_date,
    markdown_table,
)


def parse(filename: str, content: bytes, depth: int = 0) -> ParsedDocument:
    """Parse a .pptx: each slide's title, text, tables and speaker notes.

    Raises:
        EncryptedDocumentError: If the presentation is password-protected.
        ArchiveTooLargeError: If the archive fails the zip-bomb checks.
        ParseError: If python-pptx cannot read it.
    """
    open_zip(filename, content)
    try:
        from pptx import Presentation

        presentation = Presentation(BytesIO(content))
        slides = [_slide(i, slide) for i, slide in enumerate(presentation.slides, start=1)]
    except Exception as e:
        raise ParseError(
            f"Could not read '{filename}' as a PowerPoint file ({type(e).__name__})."
        ) from e

    props = presentation.core_properties
    metadata = DocumentMetadata(
        source_format="pptx",
        title=clean_str(props.title),
        author=clean_str(props.author),
        created=iso_date(props.created),
        modified=iso_date(props.modified),
        page_count=len(slides),
        language=clean_str(props.language),
    )
    return ParsedDocument(text="\n\n".join(slides), metadata=metadata)


def _slide(number: int, slide: Any) -> str:
    title_shape = slide.shapes.title
    title = " ".join(title_shape.text_frame.text.split()) if title_shape is not None else ""
    # python-pptx builds a new proxy per access, so compare shapes by id.
    title_id = title_shape.shape_id if title_shape is not None else None
    blocks = [f"## Slide {number}" + (f": {title}" if title else "")]

    for shape in slide.shapes:
        if shape.shape_id == title_id:
            continue
        if shape.has_text_frame:
            lines = [
                "  " * p.level + "- " + " ".join(p.text.split())
                for p in shape.text_frame.paragraphs
                if p.text.strip()
            ]
            blocks.append("\n".join(lines))
        elif shape.has_table:
            rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
            blocks.append(markdown_table(rows))

    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame
        if notes is not None and notes.text.strip():
            blocks.append("Notes: " + notes.text.strip())
    return "\n\n".join(b for b in blocks if b.strip())
